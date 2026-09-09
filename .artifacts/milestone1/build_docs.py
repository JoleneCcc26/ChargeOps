from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"E:\ChargeOps")
OUT = ROOT / "deliverables"
SCREENSHOT = ROOT / ".artifacts" / "milestone1" / "screenshots" / "dashboard.png"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = "14213D"
BLUE = "2E74B5"
DEEP_BLUE = "1F4D78"
TEAL = "0F766E"
PALE_BLUE = "E8F1F8"
PALE_TEAL = "E7F6F2"
PALE_GRAY = "F2F4F7"
MID_GRAY = "667085"
DARK = "17202A"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, total=9360, indent=0):
    assert sum(widths) == total, (widths, sum(widths), total)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    r_begin = OxmlElement("w:r")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    r_begin.append(fld_begin)
    r_instr = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    r_instr.append(instr)
    r_sep = OxmlElement("w:r")
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    r_sep.append(fld_sep)
    r_text = OxmlElement("w:r")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    r_text.append(fld_text)
    r_end = OxmlElement("w:r")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r_end.append(fld_end)
    paragraph._p.append(r_begin)
    paragraph._p.append(r_instr)
    paragraph._p.append(r_sep)
    paragraph._p.append(r_text)
    paragraph._p.append(r_end)


def configure_page(doc, compact=False):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.72 if compact else 0.8)
        section.bottom_margin = Inches(0.7 if compact else 0.75)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.32)
        section.footer_distance = Inches(0.32)
        header = section.header
        hp = header.paragraphs[0]
        hp.text = "CHARGEOPS  /  MILESTONE 1"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in hp.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(8)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(MID_GRAY)
        add_page_field(section.footer.paragraphs[0])


def set_paragraph_border(paragraph, side="bottom", color=BLUE, size=8, space=4):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), str(space))
    edge.set(qn("w:color"), color)
    p_bdr.append(edge)


def add_custom_numbering(doc, fmt="bullet", glyph="•"):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids + [0]) + 1
    num_id = max(num_ids + [0]) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), glyph if fmt == "bullet" else "%1.")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    lvl.append(p_pr)
    abstract.append(lvl)
    # WordprocessingML requires all abstract numbering definitions to appear
    # before concrete <w:num> instances. Keep that schema order so Word does
    # not repair the package and accidentally merge list identities.
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(list(numbering).index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_id = OxmlElement("w:abstractNumId")
    abs_id.set(qn("w:val"), str(abstract_id))
    num.append(abs_id)
    numbering.append(num)
    return num_id


def add_list_item(doc, text, num_id, bold_prefix=None):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    p_pr.append(num_pr)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def set_base_styles(doc, compact=False):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25 if compact else 1.10
    normal.paragraph_format.widow_control = True

    h1 = styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(BLUE)
    h1.paragraph_format.space_before = Pt(18 if compact else 16)
    h1.paragraph_format.space_after = Pt(10 if compact else 8)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(DEEP_BLUE)
    h2.paragraph_format.space_before = Pt(14 if compact else 12)
    h2.paragraph_format.space_after = Pt(7 if compact else 6)
    h2.paragraph_format.keep_with_next = True

    h3 = styles["Heading 3"]
    h3.font.name = "Calibri"
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor.from_string(DEEP_BLUE)
    h3.paragraph_format.space_before = Pt(10 if compact else 8)
    h3.paragraph_format.space_after = Pt(5 if compact else 4)
    h3.paragraph_format.keep_with_next = True

    if "Caption" in styles:
        cap = styles["Caption"]
        cap.font.name = "Calibri"
        cap.font.size = Pt(9)
        cap.font.italic = True
        cap.font.color.rgb = RGBColor.from_string(MID_GRAY)
        cap.paragraph_format.space_after = Pt(8)


def add_title_block(doc, title, subtitle, meta_lines, compact=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(22 if compact else 72)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(title)
    r.font.name = "Calibri"
    r.font.size = Pt(28 if compact else 31)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    set_paragraph_border(p, color=TEAL, size=18, space=10)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(18)
    r = p2.add_run(subtitle)
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    for line in meta_lines:
        p3 = doc.add_paragraph()
        p3.paragraph_format.space_after = Pt(2)
        r = p3.add_run(line)
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor.from_string(MID_GRAY)


def add_callout(doc, label, text, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label.upper() + "  ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, PALE_GRAY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string(NAVY)
    for row_data in rows:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            p = row.cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            r.font.size = Pt(9)
    return table


def add_stage_direction(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(5)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE_TEAL)
    p_pr.append(shd)
    r = p.add_run("ON SCREEN  ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)
    r2 = p.add_run(text)
    r2.italic = True
    r2.font.color.rgb = RGBColor.from_string(DEEP_BLUE)


def add_spoken(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.add_run(text)
    return p


def build_script():
    doc = Document()
    configure_page(doc, compact=True)
    set_base_styles(doc, compact=True)
    bullets = add_custom_numbering(doc, "bullet", "•")

    add_title_block(
        doc,
        "ChargeOps Milestone 1",
        "10-Minute English Video Demo Script",
        [
            "Presenter(s): [Team Member Name(s)]",
            "Course / Section: [Course / Section]",
            "Recording date: [Submission Date]",
        ],
        compact=True,
    )
    add_callout(
        doc,
        "Purpose",
        "Demonstrate a complete local application—from browser interaction, through application logic, to persistent data—while explaining the business domain and the project’s cloud-ready technical dimensions.",
    )
    doc.add_heading("Before recording", level=1)
    for item in [
        "Start the full local stack with npm run dev:all and open http://localhost:5173.",
        "Sign in with the Operations Manager demo account: ops / chargeops-demo.",
        "Keep the browser at a readable zoom level and close unrelated tabs or notifications.",
        "Have one sample technician photo or PDF ready if you plan to demonstrate Field uploads.",
        "Turn on the camera for every presenting team member during that person’s section.",
        "Rehearse once and target 9:30–10:00 so the final recording stays within the limit.",
    ]:
        add_list_item(doc, item, bullets)

    doc.add_heading("Run of show", level=1)
    add_table(
        doc,
        ["Time", "Screen / action", "Message"],
        [
            ("0:00–0:50", "Camera + title screen", "Domain, problem, and project vision"),
            ("0:50–1:30", "Architecture orientation", "Three local components working together"),
            ("1:30–2:35", "Login + Dashboard", "Role-based access and operational overview"),
            ("2:35–3:30", "Stations + Chargers", "Relational asset hierarchy and availability"),
            ("3:30–4:55", "Maintenance", "Dispatch and resolve an incident"),
            ("4:55–6:05", "Sessions + Payments", "Charging, billing, wallet, and invoices"),
            ("6:05–7:25", "Field uploads", "Unstructured files processed asynchronously"),
            ("7:25–8:35", "Cloud ops", "Queues, workers, retries, and traffic spikes"),
            ("8:35–9:20", "Daily simulation", "Realistic local platform activity"),
            ("9:20–10:00", "Camera + closing", "Milestone fit and cloud direction"),
        ],
        [1200, 3000, 5160],
    )

    doc.add_page_break()
    doc.add_heading("0:00–0:50 — Opening: the domain and the problem", level=1)
    add_stage_direction(doc, "Camera on. Show the ChargeOps title or login screen. Keep your face visible while you introduce the project.")
    add_spoken(doc, "Hello, my name is [Team Member Name], and this is ChargeOps, our EV charging network operations platform. ChargeOps is designed for the people who keep a charging network running: operations managers, field technicians, and read-only stakeholders.")
    add_spoken(doc, "The business problem is that charger status, maintenance incidents, charging sessions, payments, and field reports are closely connected, but they are often handled in separate tools. That separation makes it harder to see which charger needs attention, who is responsible, whether a charging session was billed correctly, and whether a station is ready for customers.")
    add_spoken(doc, "We chose this domain because it has a clear operational workflow and meaningful technical depth. It requires structured relational data, unstructured technician files, and asynchronous processing for billing, telemetry, and traffic spikes. Our goal is to give an internal operations team one reliable view of the network.")

    doc.add_heading("0:50–1:30 — Local architecture", level=1)
    add_stage_direction(doc, "Briefly show Slide 2, the README architecture section, or a simple architecture diagram. Then return to the application.")
    add_spoken(doc, "For Milestone 1, every component runs locally. The front end is a React and Vite web application. The application layer is a Node.js and Express API with authentication, role-based authorization, queue adapters, and background workers. The primary data store is MySQL. The application also uses local object storage for photos, PDFs, and generated invoice files.")
    add_spoken(doc, "The browser never talks directly to the database. A user action calls the API, the API validates authorization and business rules, and then it reads or writes persistent data. Longer-running work is placed on a queue and completed by workers. That separation gives us a complete local system now and clear seams for later cloud migration.")

    doc.add_heading("1:30–2:35 — Login and dashboard", level=1)
    add_stage_direction(doc, "Sign in as ops using ops / chargeops-demo. Pause on the Dashboard and point to the KPI cards, charts, and recent activity.")
    add_spoken(doc, "I will sign in as the Operations Manager. Authentication happens through the application layer, and the returned session determines which routes and actions this user can access. We also have a technician role for field work and a viewer role for read-only access.")
    add_spoken(doc, "The dashboard summarizes live operational data from MySQL. It shows the station and charger inventory, the user base, plan information, charger availability, revenue by city, and recent activity. These values are not hard-coded cards; the front end requests them from the API, and the API aggregates the underlying relational tables.")
    add_spoken(doc, "This page gives an operations manager a fast answer to three questions: how large is the network, how much of it is available, and where should the team investigate next?")

    doc.add_heading("2:35–3:30 — Stations and chargers", level=1)
    add_stage_direction(doc, "Open Stations, search or filter once, and select a station. Then open Chargers and point out status, type, and station association.")
    add_spoken(doc, "Stations and chargers demonstrate the asset hierarchy in our relational model. A company operates stations, and each station contains one or more chargers. The interface lets the operator search the network, inspect location and availability, and connect a physical charger to the sessions and incidents recorded against it.")
    add_spoken(doc, "The important point is consistency across the system. A charger displayed here is the same charger referenced by a charging session, a payment, and a maintenance record. Foreign-key relationships prevent the application from creating disconnected operational records.")

    doc.add_heading("3:30–4:55 — Core maintenance workflow", level=1)
    add_stage_direction(doc, "Open Maintenance. Choose an open item. Assign a technician from the dropdown, click Start, and then Resolve if you are comfortable changing the demo record. Show the resulting status after each action.")
    add_spoken(doc, "Maintenance is one of our core user flows. An open incident identifies the charger, station, reported issue, assigned technician, status, and resolution time. As an operations manager, I can assign the correct technician. That change goes through an authorized API endpoint and is persisted in the maintenance log.")
    add_spoken(doc, "Next, I can start the work. The status changes from Open to In Progress. When the issue is complete, I can resolve it, and the system records the resolved time. These buttons are not only visual state changes: each transition is validated by the application layer and saved in MySQL.")
    add_spoken(doc, "Role-based controls matter here. An operations manager can coordinate the entire network. A technician sees the workflow needed for field service. A viewer can inspect information without changing it. This makes the product primarily an internal operations application, not a public driver-facing charging app.")

    doc.add_heading("4:55–6:05 — Sessions, payments, wallets, and invoices", level=1)
    add_stage_direction(doc, "Open Sessions and select a completed example. Then open Payments and show its matching payment or invoice metadata.")
    add_spoken(doc, "A charging session connects a user, vehicle, charger, start and end time, energy consumption, and total cost. When a completed session is billed, the application creates a payment and updates the user wallet in one controlled workflow. Successful billing also produces invoice metadata and a PDF file.")
    add_spoken(doc, "This is where transactional integrity is important. We validate that a session is billed at most once, that payment status matches the session outcome, and that wallet balances do not become negative. We also repaired legacy demo data before this milestone by releasing stuck chargers, expiring stale subscriptions, cancelling completed-but-unpaid sessions, failing an abandoned pending payment, and restoring missing invoice metadata.")
    add_spoken(doc, "The final result is traceable from physical charging activity to a financial record. The operations team can investigate both service reliability and revenue without leaving the platform.")

    doc.add_heading("6:05–7:25 — Field uploads and unstructured files", level=1)
    add_stage_direction(doc, "Open Field uploads. Upload a prepared photo or PDF, or show an already processed upload if recording time is limited. Point out the asynchronous status and linked maintenance record.")
    add_spoken(doc, "The relational database is not the only data source. Field technicians also upload photos and PDF reports. The API stores the file in local object storage, creates metadata, places a file-processing job on a queue, and immediately returns an accepted response instead of blocking the browser.")
    add_spoken(doc, "A background worker then parses the file. For images, it can inspect JPEG and EXIF metadata. For PDFs, it extracts text. The worker classifies the likely fault type, severity, and error code, compares GPS information with station locations when available, and then creates or updates the corresponding maintenance ticket.")
    add_spoken(doc, "This demonstrates our unstructured-file dimension and application logic beyond basic create, read, update, and delete screens. A technician’s evidence becomes a searchable, actionable operational record.")

    doc.add_heading("7:25–8:35 — Asynchronous operations and traffic spikes", level=1)
    add_stage_direction(doc, "Open Cloud ops. Show queue depth, worker count, throughput or recent jobs. If the system is idle, explain that zero backlog is the healthy state.")
    add_spoken(doc, "Cloud ops exposes the asynchronous part of the platform. Billing, file processing, telemetry, and simulation jobs are separated into queues. Workers claim jobs, update their progress, and complete them outside the request-response path.")
    add_spoken(doc, "The worker system includes retry delays, dead-letter handling for repeated failures, and idempotency protections so the same request does not create duplicate business records. A local autoscaling supervisor can increase or reduce worker capacity based on backlog. This is especially relevant to EV charging because device telemetry and session completion events can arrive in bursts.")
    add_spoken(doc, "For Milestone 1, the queue is implemented locally with MySQL-backed job records. The design is intentionally replaceable, so a later cloud version can map the same boundary to a managed queue and horizontally scaled workers.")

    doc.add_heading("8:35–9:20 — Daily operating simulation", level=1)
    add_stage_direction(doc, "Show the simulation status in Cloud ops, or briefly show the terminal command npm run simulate:status. Do not trigger a second run if today is already complete.")
    add_spoken(doc, "To keep the demo environment representative of a normally operating platform, ChargeOps includes a daily simulator. A MySQL event schedules one run per date, and a worker generates realistic charging sessions, payments, invoices, telemetry, and maintenance situations.")
    add_spoken(doc, "The current run generated eighteen sessions and two maintenance situations. All eighteen payments and invoices completed successfully, with no negative wallets or duplicate billing. The simulator is bounded and idempotent, so it creates useful activity without continuously corrupting the data.")

    doc.add_heading("9:20–10:00 — Closing", level=1)
    add_stage_direction(doc, "Return to camera. Optionally keep the dashboard visible beside your camera feed.")
    add_spoken(doc, "To summarize, ChargeOps is a working local application with all three required components. The React front end supports real operator workflows. The Express application layer enforces authentication, roles, business rules, and asynchronous jobs. MySQL and local object storage persist both structured and unstructured data.")
    add_spoken(doc, "We are proud of this domain because the user flow is easy to understand while the architecture addresses real cloud-computing concerns: relational consistency, file processing, and traffic spikes. We have also verified the project with a successful production build, twenty-six extraction self-tests, and twenty-six authenticated API smoke tests.")
    add_spoken(doc, "In the next milestone, these local seams can move to managed cloud services for web delivery, containers, relational storage, object storage, and queues. Thank you for watching our ChargeOps Milestone 1 demonstration.")

    doc.add_heading("Presenter handoff and evidence checklist", level=1)
    for item in [
        "Replace every bracketed placeholder before recording or submission.",
        "Keep each presenting team member’s face visible during that person’s spoken portion.",
        "Show at least one write operation that visibly persists after refresh.",
        "Show the front end, explain the API/application logic, and identify MySQL plus object storage.",
    ]:
        add_list_item(doc, item, bullets)

    path = OUT / "ChargeOps_Milestone1_Video_Script.docx"
    doc.save(path)
    return path


def build_report():
    doc = Document()
    configure_page(doc, compact=False)
    set_base_styles(doc, compact=False)
    bullets = add_custom_numbering(doc, "bullet", "•")

    add_title_block(
        doc,
        "ChargeOps",
        "Milestone 1 Local Application Report",
        [
            "EV Charging Network Operations Platform",
            "Team: [Team Member Name(s)]",
            "Course / Section: [Course / Section]",
            "Submission date: [Submission Date]",
        ],
        compact=False,
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(54)
    add_callout(
        doc,
        "Milestone outcome",
        "ChargeOps is a fully functional local application with an integrated React front end, Node.js/Express application layer, MySQL relational store, local object storage, and background job workers.",
        fill=PALE_TEAL,
    )
    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph(
        "ChargeOps gives an internal EV charging operations team one place to monitor infrastructure, coordinate field maintenance, review charging activity, and reconcile payments. The Milestone 1 implementation demonstrates the complete user path from a browser action to application logic and persistent data. It also establishes local equivalents of cloud-oriented capabilities—object storage, queues, workers, telemetry ingestion, and autoscaling supervision—without claiming that the system has already migrated to the cloud."
    )

    doc.add_page_break()
    doc.add_heading("1. Business domain and project vision", level=1)
    doc.add_heading("Problem statement", level=2)
    doc.add_paragraph(
        "An EV charging network is both a physical service network and a transaction platform. Operators must understand charger availability, associate incidents with the correct equipment and station, dispatch technicians, confirm that completed sessions are billed once, and retain evidence such as field photos and service reports. When these activities are isolated in separate tools, incident response slows down and operational data becomes difficult to reconcile."
    )
    doc.add_heading("Target users", level=2)
    add_table(
        doc,
        ["Role", "Primary responsibility", "Typical actions"],
        [
            ("Operations Manager", "Network-wide coordination", "Monitor KPIs; assign technicians; start or resolve work; investigate sessions and payments"),
            ("Field Technician", "On-site service execution", "Review assigned work; submit photos/PDF reports; update service progress"),
            ("Read-only Viewer", "Operational visibility", "Inspect dashboards and records without changing production state"),
        ],
        [1900, 2850, 4610],
    )
    doc.add_heading("Why this domain", level=2)
    doc.add_paragraph(
        "The domain was selected because it supports a clear core workflow while naturally exercising three cloud-computing dimensions. Relational integrity is required across users, chargers, sessions, payments, wallets, and maintenance records. Unstructured photos, PDF service reports, and invoice files require durable object handling. Telemetry, uploads, billing, and end-of-session events benefit from asynchronous processing when traffic arrives in bursts."
    )
    doc.add_heading("Project vision", level=2)
    add_callout(
        doc,
        "Vision",
        "Give the operations team a trustworthy path from network condition to field action to financial outcome, with every important change traceable through the application and data layers.",
    )

    doc.add_page_break()
    doc.add_heading("2. What the platform does", level=1)
    doc.add_paragraph(
        "ChargeOps is an internal operations console, not a consumer app. Drivers never see it. "
        "Its users are the people responsible for keeping a charging network available: the operations "
        "manager who owns uptime across every site, and the field technicians they dispatch. "
        "The platform exists to close two loops that are usually spread across disconnected tools — "
        "turning a broken charger back into a working one, and turning a completed charging session "
        "into a paid, receipted transaction."
    )

    doc.add_heading("What each role can do", level=2)
    doc.add_paragraph(
        "Permissions are enforced by the API, not by the interface. The front end hides menu items a role "
        "cannot use, but that is a convenience: a technician who calls a restricted endpoint directly still "
        "receives HTTP 403. Work orders are additionally filtered row by row in SQL, so a technician "
        "retrieves only the jobs assigned to them."
    )
    add_table(
        doc,
        ["Capability", "Operations Manager", "Field Technician", "Read-only Viewer"],
        [
            ("Network dashboard, stations, chargers", "Yes", "Yes", "Yes"),
            ("Work orders — view", "Entire network", "Own assignments only", "Entire network"),
            ("Work orders — assign to a technician", "Yes", "No", "No"),
            ("Work orders — change status / resolve", "Any work order", "Own assignments only", "No"),
            ("Reopen a resolved work order", "Yes", "No", "No"),
            ("Upload field photos and service reports", "Yes", "Yes", "No"),
            ("Customer directory, sessions, payments, invoices", "Yes", "No", "Yes"),
            ("Queue and worker monitoring", "Yes", "No", "Yes"),
            ("Administrative queue actions (purge, redrive)", "Yes", "No", "No"),
        ],
        [3400, 2100, 2100, 1760],
    )

    doc.add_heading("Who signs in", level=2)
    doc.add_paragraph(
        "Work is assigned to individual technicians, and the network employs about a hundred of them, so "
        "technician identity comes from the business data rather than from a fixed list of accounts. Any "
        "technician on record signs in with their own identifier and sees their own queue. This is what "
        "makes the dispatch handoff demonstrable: the manager assigns a job to whoever is closest to the "
        "site, and that specific person can then sign in and pick it up."
    )
    add_table(
        doc,
        ["Account", "Resolves to", "Sees"],
        [
            ("Operations manager", "A fixed operations account", "The whole network, plus queue and worker monitoring"),
            ("Technician", "A row in the technician table, by identifier", "Only the work orders assigned to that person"),
            ("Read-only viewer", "A fixed read-only account", "The whole network, with every write action refused"),
        ],
        [2400, 3200, 3760],
    )
    add_callout(
        doc,
        "Known limitation",
        "Demo technician accounts share a single password held in plain text, and the token provider signs "
        "locally with a shared secret. Both are acceptable for a local milestone and neither is acceptable "
        "in production: a real deployment needs per-user password hashing or an external identity provider. "
        "The change is contained in one function and does not affect the authorization model, which is "
        "already enforced per request and per row.",
    )

    doc.add_heading("Loop 1: a broken charger becomes a working charger", level=2)
    doc.add_paragraph(
        "A charger can be reported faulty two ways: it reports the fault itself over telemetry, or a "
        "technician uploads a service report that the system classifies as critical. Either path produces "
        "the same thing — an open work order — and that work order is what makes the problem visible "
        "and assignable."
    )
    loop1 = add_custom_numbering(doc, "decimal", "%1.")
    for item in [
        "A fault is detected. The charger reports a fault code, or a technician's uploaded report is classified as critical.",
        "The charger is taken out of service, so no driver is sent to a stall that cannot charge them.",
        "A work order is opened automatically and pre-assigned to a technician near that station.",
        "The operations manager reviews the queue and assigns or reassigns the work to the right technician.",
        "The technician sees the job in their own queue, attends the site, uploads photo evidence, and marks the work in progress.",
        "The technician resolves the work order. The charger returns to service and its last-maintenance date is stamped automatically.",
    ]:
        add_list_item(doc, item, loop1)
    add_callout(
        doc,
        "Why this matters",
        "A charger is out of service because a work order is open, and resolving that work order is the only "
        "thing that brings it back. Because the rule holds in both directions, a charger can never be "
        "stranded offline with nobody responsible for it.",
    )

    add_callout(
        doc,
        "A naming correction",
        "In the inherited schema a session that was actively delivering power was stored as 'Pending'. That "
        "word means waiting, so the state read as a session that had not started, when it was in fact the "
        "most active thing in the system - and because energy and cost are only known once a session ends, "
        "those rows displayed as 0 kWh and $0.00, reinforcing the impression that nothing had happened. The "
        "state is now called 'Active', the lifecycle reads Active to Completed or Cancelled, and the "
        "interface shows elapsed charging time in place of a zero. The subscription table keeps its own "
        "'Pending' state, where the word is accurate: a subscription whose start date is in the future "
        "genuinely has not begun.",
    )

    doc.add_heading("Loop 2: a charging session becomes a paid invoice", level=2)
    loop2 = add_custom_numbering(doc, "decimal", "%1.")
    for item in [
        "A session starts. The charger is marked in use and a session record is opened.",
        "The driver unplugs. The application records the end time, frees the charger immediately, and accepts the request without calculating anything.",
        "A billing worker computes energy delivered, applies any active membership discount, and takes payment from the wallet or the card on file.",
        "A second worker renders a PDF invoice, stores the file, and records its location against the session.",
        "The operations manager can trace any completed session to its payment, wallet movement, and downloadable invoice.",
    ]:
        add_list_item(doc, item, loop2)
    doc.add_paragraph(
        "Steps 3 and 4 happen after the driver's request has already been answered. That separation is "
        "deliberate and is described in section 5."
    )

    doc.add_page_break()
    doc.add_heading("3. Core user flow", level=1)
    doc.add_paragraph(
        "The primary Milestone 1 demonstration follows an operations manager from authentication through incident resolution and billing review. Each step exercises the front end, API/application logic, and persistent store."
    )
    numbered = add_custom_numbering(doc, "decimal", "%1.")
    flow_items = [
        "Sign in. The user authenticates through the API and receives permissions associated with an operations role.",
        "Review network health. The dashboard requests aggregated counts, charger availability, revenue by city, and recent operational activity.",
        "Inspect assets. Station and charger views expose the relational hierarchy connecting locations, equipment, sessions, and incidents.",
        "Dispatch maintenance. The manager assigns a technician, starts work, and resolves the incident; the application validates each transition and persists status and timestamps.",
        "Trace charging revenue. The manager connects a completed charging session to payment status, wallet movement, invoice metadata, and the invoice file.",
        "Process field evidence. A technician uploads a photo or PDF; the API stores the file, queues processing, and a worker classifies the report and creates or updates the maintenance record.",
        "Observe asynchronous health. Cloud ops shows queue depth, worker activity, retry/dead-letter behavior, and locally supervised capacity.",
    ]
    for item in flow_items:
        add_list_item(doc, item, numbered)

    if SCREENSHOT.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        shape = run.add_picture(str(SCREENSHOT), width=Inches(6.45))
        shape._inline.docPr.set("descr", "ChargeOps local operations dashboard with station, charger, user, availability, and revenue metrics")
        cap = doc.add_paragraph("Figure 1. ChargeOps local dashboard populated from the application API and MySQL data.", style="Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()
    doc.add_heading("4. Local architecture", level=1)
    doc.add_paragraph(
        "ChargeOps uses a layered architecture. The components run together on the local machine but communicate through stable interfaces that can be replaced with managed cloud services later."
    )
    add_table(
        doc,
        ["Component", "Milestone 1 implementation", "Responsibility"],
        [
            ("Front end", "React 18, Vite, Tailwind CSS, Recharts, Leaflet", "Role-based operations interface; dashboards, maps, searchable tables, and workflow controls"),
            ("Application layer", "Node.js, Express, REST API, authentication/RBAC, queue adapters, workers", "Validate requests; enforce business rules; orchestrate billing, uploads, telemetry, and maintenance"),
            ("Data store", "MySQL plus local object storage", "Persist relational records, job state, telemetry metadata, photos, PDFs, and generated invoices"),
        ],
        [1650, 3320, 4390],
    )
    doc.add_heading("How a request travels through the three layers", level=2)
    doc.add_paragraph(
        "Requests take one of two shapes, and choosing between them is the central design decision of the "
        "application layer."
    )
    doc.add_paragraph(
        "Short work is done inside the request. Assigning a work order is one authenticated call: Express "
        "verifies the token and the role, opens a transaction, locks the work-order row, updates it, returns "
        "the charger to service if the work is resolved, writes an audit entry, and commits. The manager sees "
        "the result in roughly ten milliseconds. Nothing is queued, because the manager is waiting and the "
        "work is two row updates.",
    )
    doc.add_paragraph(
        "Long work is accepted and deferred. When a driver ends a charging session the application records "
        "the end time, frees the charger, writes a job record, and answers immediately with HTTP 202 "
        "Accepted. A background worker then computes the bill, moves the money, and produces the invoice "
        "file. Measured locally, the accepted-and-deferred path answers in about twelve milliseconds at the "
        "median; computing the same bill inside the request takes over four hundred. The work is identical, "
        "but nobody waits for it."
    )
    add_callout(
        doc,
        "The rule",
        "Queue the work that is slow, deferrable, and safe to retry. Do everything else in the request. "
        "Putting fast interactive operations on a queue would add latency and complexity for no benefit.",
    )

    doc.add_heading("Relational model", level=2)
    doc.add_paragraph(
        "The principal entities are company, user, company membership, technician, station, charger, wallet, subscription, charging session, payment, and maintenance log. Their relationships allow the application to trace one operational event across physical infrastructure, user activity, and financial records."
    )
    doc.add_heading("How unstructured files are stored", level=2)
    doc.add_paragraph(
        "Photographs and PDF reports are never written into MySQL. The file bytes go to object storage, and "
        "the database records only where the file is and what was learned from it. The two stores answer "
        "different questions: a relational database is built for small rows that are joined and updated "
        "constantly, while an object store is built for large immutable files that are written once and read "
        "occasionally. Storing files as database blobs displaces real query data from memory and makes every "
        "backup grow with the media volume."
    )
    file_steps = add_custom_numbering(doc, "decimal", "%1.")
    for item in [
        "The technician uploads a photo or report. The application writes the bytes to object storage under a dated key, records a metadata row marked pending, queues a processing job, and returns immediately.",
        "A file worker reads the object back and extracts what it can: image dimensions from the file header, capture time and GPS coordinates from the photograph's embedded metadata, and text from the PDF.",
        "The extracted text is classified into a fault category, a severity, a vendor error code, and referenced part numbers. A photograph's GPS position is matched against known station coordinates, so the image identifies the site it was taken at.",
        "Those results are written back to the metadata row as ordinary columns, which makes them searchable and reportable, and the file is attached to a work order.",
        "When the interface needs to display the file, the server issues a signed link that expires after fifteen minutes. The bucket itself is never publicly readable.",
    ]:
        add_list_item(doc, item, file_steps)
    doc.add_paragraph(
        "This is the step that converts unstructured evidence into operational data: a photograph and a "
        "paragraph of prose become a categorised, severity-ranked, site-linked work order without anyone "
        "completing a form."
    )

    doc.add_heading("Authorization boundary", level=2)
    doc.add_paragraph(
        "The front end may hide unavailable actions for usability, but the API remains the authoritative security boundary. It checks authentication and role permissions before allowing assignments, status changes, billing operations, or administrative access."
    )

    doc.add_page_break()
    doc.add_heading("5. Technical dimensions", level=1)
    add_table(
        doc,
        ["Dimension", "ChargeOps evidence", "Why it matters for later cloud work"],
        [
            ("Relational data", "Stations, chargers, users, sessions, payments, wallets, subscriptions, maintenance, foreign keys, and transactional billing", "Managed relational storage, backups, constraints, indexing, and transaction reliability"),
            ("Unstructured files", "Technician JPEG/PDF uploads and generated invoice PDFs in local object storage", "Durable object storage, lifecycle rules, controlled access, and event-driven processing"),
            ("Async / traffic spikes", "Separate queues for billing, files, telemetry, and simulation; worker pool; retries, backoff, dead-letter handling, and idempotency", "Decoupled managed queues and horizontally scalable consumers during bursty demand"),
        ],
        [1800, 3920, 3640],
    )
    doc.add_heading("Unstructured-file workflow", level=2)
    doc.add_paragraph(
        "When a technician uploads a field photo or PDF, the API writes the binary file to local object storage and records metadata. A file worker parses image metadata or PDF text, classifies fault type, severity, and error code, and can compare GPS information with known station locations. The result is linked to a maintenance ticket, turning field evidence into an operational action."
    )
    doc.add_heading("Asynchronous reliability", level=2)
    doc.add_paragraph(
        "Queue jobs carry persistent status so a transient process restart does not erase the work. Retry delays and dead-letter state make failures visible instead of creating an endless loop. Idempotency prevents a repeated request from generating duplicate payments, invoice records, or simulation runs. The API can acknowledge accepted work quickly while workers complete the expensive step independently."
    )
    doc.add_heading("Cloud migration seams", level=2)
    doc.add_paragraph(
        "The current local design does not depend on a specific cloud provider. Conceptually, the static web client can move behind a content-delivery service; the Express API and workers can run as containers; MySQL can move to a managed relational database; object files can move to object storage; and MySQL-backed queues can be replaced with a managed queue. These are future mappings, not requirements for Milestone 1."
    )

    doc.add_page_break()
    doc.add_heading("6. Operational realism and data integrity", level=1)
    doc.add_heading("Daily simulation", level=2)
    doc.add_paragraph(
        "A scheduled MySQL event enqueues one bounded simulation per date. The worker creates realistic platform activity: charging sessions, charger telemetry, payments, invoice files, and maintenance situations. A run typically produces twelve to twenty-four sessions and one to three maintenance situations, then resolves a subset of older simulated incidents to keep the environment active without unbounded growth."
    )
    doc.add_paragraph(
        "The verified August 31, 2026 run produced 18 sessions and 2 situations, totaling 1,807.80 kWh and $792.96. All 18 payments and invoice records completed successfully. Integrity checks found no negative wallets and no duplicate billing."
    )
    doc.add_heading("Keeping the seeded data meaningful over time", level=2)
    doc.add_paragraph(
        "The seed data carried over from the database project describes a fixed window of dates. Wall-clock "
        "time keeps moving and those dates do not, so parts of the dataset silently stop being usable. The "
        "clearest example: every membership subscription in the original data had lapsed, which meant no "
        "driver held an active membership and the billing worker's discount calculation could never run. A "
        "whole feature had quietly disappeared from the application without anything failing."
    )
    doc.add_paragraph(
        "Setup therefore re-anchors subscriptions to the current date each time it runs, producing a "
        "deliberately mixed population rather than making every membership active: roughly seventy percent "
        "current, twenty percent lapsed, ten percent starting shortly. Placement is derived from the "
        "subscription identifier rather than randomly, so every teammate's database and every recording of "
        "the demo shows the same figures. Cancelled subscriptions are left alone, because cancellation is a "
        "customer decision rather than a date that drifted."
    )

    doc.add_heading("Making the data look like a network that is running", level=2)
    doc.add_paragraph(
        "Correct data is not the same as convincing data. A charging network is never idle: at any moment "
        "some chargers are mid-session, yesterday resembled today, and demand rises and falls with commuting "
        "hours. Early in testing the database satisfied every integrity check and still read as an empty "
        "car park, because the load generator opened sessions and closed every one of them, leaving no "
        "charger in use, and wrote its thousands of records into whichever clock hour the test happened to "
        "run in."
    )
    doc.add_paragraph(
        "A load generator is a performance tool and was never meant to double as business history, so shaping "
        "the operating picture is a separate step. It spreads bulk-written sessions across the preceding "
        "weeks along a realistic daily demand curve with morning and evening peaks, moving each payment and "
        "invoice with its session so the money still lines up. It then opens genuine sessions on roughly a "
        "fifth of the fleet, spread across every station and across both charger types, with start times "
        "staggered over the previous two hours."
    )
    doc.add_paragraph(
        "Those sessions are real rows rather than display values, which is what makes them useful: the "
        "operations manager can end one during the demonstration and watch the billing pipeline run against "
        "it. Placement is derived from record identifiers rather than randomly, so every teammate's database "
        "and every recording of the demonstration show the same figures."
    )

    doc.add_heading("Data integrity: causes fixed, not just symptoms", level=2)
    doc.add_paragraph(
        "Two classes of inconsistency appeared during testing. Both were traced to their cause and corrected "
        "in the application rather than cleaned up repeatedly in the database."
    )
    for item in [
        "Chargers stranded out of service. Telemetry could take a charger out of service without raising a work order, so nothing appeared in the maintenance queue, no technician was ever dispatched, and the charger had no route back into service. Telemetry now opens a work order whenever it takes a charger down, which closes the loop: fault, work order, dispatch, resolution, back in service.",
        "Chargers left occupied by the load generator. The load profile reported randomly chosen statuses, which rewrote charger availability as a side effect of a performance test. It now reports each charger's actual state, so the load test stresses the ingest pipeline without altering operational data.",
    ]:
        add_list_item(doc, item, bullets)
    doc.add_paragraph(
        "An audit command reports every known class of inconsistency read-only, and a repair command applies "
        "the corrections inside a single transaction after writing a rollback snapshot, verifying the "
        "expected end state, and recording an audit entry. It exists as a safety net; with the causes fixed, "
        "a full load test now leaves the audit reporting zero findings across every category."
    )
    doc.add_heading("Verification evidence", level=2)
    add_table(
        doc,
        ["Check", "Result", "What it demonstrates"],
        [
            ("Production front-end build", "Passed", "The React application compiles for deployment."),
            ("Extraction self-test", "26 passed", "EXIF parsing, PDF text extraction, fault classification, and GPS-to-station matching, run without a database."),
            ("Authenticated API smoke test", "26 passed", "Login, role permissions, per-technician work-order scoping, pagination, and the main endpoints against a live server and database."),
            ("Simulation integrity", "Passed", "No negative wallets or duplicate billing in the verified run."),
        ],
        [2600, 1500, 5260],
    )

    doc.add_page_break()
    doc.add_heading("7. Milestone 1 requirement alignment", level=1)
    add_table(
        doc,
        ["Requirement / rubric area", "ChargeOps submission evidence"],
        [
            ("Working demo — functionality", "Login, dashboard, station/charger inspection, maintenance transitions, sessions/payments, file uploads, and queue monitoring run locally end to end."),
            ("Domain fit and technical justification", "EV charging operations has a clear internal user, business workflow, relational model, unstructured files, and bursty asynchronous workloads."),
            ("Video demo quality", "The companion script allocates ten minutes, includes camera cues, and walks through a coherent core user flow instead of disconnected screens."),
            ("PPT clarity — maximum two slides", "Slide 1 states the name, problem, users, team placeholders, and domain rationale. Slide 2 shows the three-layer architecture and technical dimensions."),
            ("Team collaboration evidence", "Replace the placeholders below with team roles, commits, meeting notes, or a shared task record before submission."),
        ],
        [3150, 6210],
    )
    doc.add_heading("Team collaboration record", level=2)
    add_table(
        doc,
        ["Team member", "Contribution", "Evidence to attach or reference"],
        [
            ("[Team Member Name]", "[Front end / application / database / testing / presentation]", "[Commit, pull request, task board item, or meeting note]"),
            ("[Team Member Name]", "[Contribution]", "[Evidence]"),
            ("[Team Member Name]", "[Contribution]", "[Evidence]"),
        ],
        [2400, 3520, 3440],
    )
    doc.add_heading("Conclusion", level=2)
    doc.add_paragraph(
        "ChargeOps meets the Milestone 1 objective by demonstrating a viable domain and a functioning local front end, application layer, and data store. The project’s strongest evidence is the connected workflow: operational records are not isolated screens, but traceable data that moves from charger activity and field evidence through validated application logic to maintenance and financial outcomes."
    )
    add_callout(
        doc,
        "Before submitting",
        "Replace all bracketed placeholders, confirm the final demo account and sample record, keep presenter cameras on, and submit this report together with the two-slide deck and video.",
        fill=PALE_TEAL,
    )

    path = OUT / "ChargeOps_Milestone1_Report.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    for created in (build_script(), build_report()):
        print(created)
